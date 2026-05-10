"""Excel export helpers for chat_ops raw report tabs."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from otutil.tools._chat_ops.common import query_rows
from otutil.tools._chat_ops.modes import run_mode

RAW_MODES: tuple[str, ...] = (
    "commands",
    "invocations",
    "files",
    "signals",
    "raw",
    "sessions",
    "annotations",
    "usage",
    "turns",
)


@dataclass(slots=True)
class ExcelRequest:
    """Parameters for report_excel exports."""

    modes: list[str]
    projects: list[str] | None
    session_ids: list[str] | None
    models: list[str] | None
    start: str | None
    end: str | None
    limit: int | None


@dataclass(slots=True)
class _SessionDim:
    project: str | None
    started_at: str | None


@dataclass(slots=True)
class _ModeReq:
    group_by: str | None = None
    project: str | None = None
    session_id: str | None = None
    model: str | None = None
    start: str | None = None
    end: str | None = None
    min_samples: int = 1
    limit: int = 100


_MODE_TABLES: dict[str, str] = {
    "commands": "event_commands",
    "invocations": "event_invocations",
    "files": "event_file_ops",
    "signals": "event_signals",
    "raw": "events",
    "sessions": "mv_session_metrics",
    "annotations": "event_annotations",
    "usage": "event_usage",
    "turns": "mv_turn_metrics",
}


def normalize_modes(value: list[str] | None) -> list[str]:
    """Normalize requested raw mode names."""
    if value is None:
        return list(RAW_MODES)
    if not isinstance(value, list):
        raise TypeError("report must be a list[str] when provided")

    raw = [str(item).strip() for item in value if str(item).strip()]

    if not raw:
        return list(RAW_MODES)

    out: list[str] = []
    seen: set[str] = set()
    for mode in raw:
        if mode not in RAW_MODES:
            raise ValueError(
                f"unsupported excel report mode '{mode}'. Allowed: {', '.join(RAW_MODES)}"
            )
        if mode not in seen:
            seen.add(mode)
            out.append(mode)
    return out


def normalize_filter_values(value: str | list[str] | None) -> list[str] | None:
    """Normalize list-friendly filter input."""
    if value is None:
        return None
    if isinstance(value, str):
        items = [part.strip() for part in value.split(",") if part.strip()]
    else:
        items = [str(part).strip() for part in value if str(part).strip()]
    return items or None


def resolve_excel_output_path(*, output_dir: Path, output_name: str | None) -> Path:
    """Resolve output path for excel report."""
    raw_name = (output_name or "").strip() or "report-excel.xlsx"
    out = Path(raw_name)
    if out.name != raw_name:
        raise ValueError("output_name must be a filename, not a path")
    if not out.stem:
        raise ValueError("output_name must be non-empty")
    suffix = out.suffix.lower()
    if suffix and suffix != ".xlsx":
        raise ValueError("output_name extension must be '.xlsx'")
    filename = raw_name if suffix else f"{raw_name}.xlsx"
    return output_dir / filename


def _parse_iso(value: str | None) -> datetime | None:
    if not value:
        return None
    text = value.strip()
    if not text:
        return None
    normalized = text.replace("Z", "+00:00")
    dt = datetime.fromisoformat(normalized)
    if dt.tzinfo is None:
        return dt.replace(tzinfo=UTC)
    return dt.astimezone(UTC)


def _load_session_dims(conn: Any, *, session_ids: set[str]) -> dict[str, _SessionDim]:
    if not session_ids:
        return {}

    placeholders = ", ".join(f":sid_{idx}" for idx, _ in enumerate(session_ids))
    params = {f"sid_{idx}": sid for idx, sid in enumerate(session_ids)}
    rows = query_rows(
        conn,
        sql=f"""
            SELECT session_id, project, started_at
            FROM mv_session_metrics
            WHERE session_id IN ({placeholders})
        """,
        params=params,
    )
    out: dict[str, _SessionDim] = {}
    for row in rows:
        sid = str(row.get("session_id") or "")
        out[sid] = _SessionDim(
            project=row.get("project"),
            started_at=row.get("started_at"),
        )
    return out


def _load_turn_models(
    conn: Any,
    *,
    session_ids: set[str],
) -> tuple[dict[tuple[str, str], str | None], dict[str, str | None]]:
    if not session_ids:
        return {}, {}

    placeholders = ", ".join(f":sid_{idx}" for idx, _ in enumerate(session_ids))
    params = {f"sid_{idx}": sid for idx, sid in enumerate(session_ids)}
    rows = query_rows(
        conn,
        sql=f"""
            SELECT session_id, COALESCE(turn_id, '') AS turn_id, model, turn_started_at
            FROM mv_turn_metrics
            WHERE session_id IN ({placeholders})
            ORDER BY session_id ASC, turn_started_at DESC
        """,
        params=params,
    )
    by_turn: dict[tuple[str, str], str | None] = {}
    latest_by_session: dict[str, str | None] = {}
    for row in rows:
        sid = str(row.get("session_id") or "")
        tid = str(row.get("turn_id") or "")
        model = row.get("model")
        key = (sid, tid)
        if key not in by_turn:
            by_turn[key] = model
        if sid not in latest_by_session:
            latest_by_session[sid] = model
    return by_turn, latest_by_session


def _canonical_date(started_at: str | None) -> str | None:
    if not started_at:
        return None
    return started_at[:10]


def _row_session_id(row: dict[str, Any]) -> str | None:
    val = row.get("session_id")
    if val is None:
        return None
    text = str(val).strip()
    return text or None


def _row_turn_id(row: dict[str, Any]) -> str:
    val = row.get("turn_id")
    if val is None:
        return ""
    return str(val)


def _keep_row(
    *,
    project: str | None,
    session_id: str | None,
    model: str | None,
    started_at: str | None,
    req: ExcelRequest,
) -> bool:
    if req.projects and (project not in req.projects):
        return False
    if req.session_ids and (session_id not in req.session_ids):
        return False
    if req.models and (model not in req.models):
        return False

    start_dt = _parse_iso(req.start)
    end_dt = _parse_iso(req.end)
    if start_dt is None and end_dt is None:
        return True

    row_dt = _parse_iso(started_at)
    if row_dt is None:
        return False
    if start_dt and row_dt < start_dt:
        return False
    return not (end_dt and row_dt > end_dt)


def _enrich_rows(
    *,
    rows: list[dict[str, Any]],
    req: ExcelRequest,
    session_dims: dict[str, _SessionDim],
    turn_models: dict[tuple[str, str], str | None],
    latest_models: dict[str, str | None],
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in rows:
        session_id = _row_session_id(row)
        turn_id = _row_turn_id(row)
        dim = session_dims.get(session_id or "")

        project = row.get("project") if row.get("project") is not None else (dim.project if dim else None)
        model = row.get("model")
        if model is None and session_id is not None:
            model = turn_models.get((session_id, turn_id))
        if model is None and session_id is not None:
            model = latest_models.get(session_id)
        started_at = dim.started_at if dim else None
        date = _canonical_date(started_at)

        if not _keep_row(
            project=project if isinstance(project, str) else None,
            session_id=session_id,
            model=model if isinstance(model, str) else None,
            started_at=started_at,
            req=req,
        ):
            continue

        canonical = {
            "project": project,
            "session_id": session_id,
            "model": model,
            "date": date,
        }
        remainder: dict[str, Any] = {}
        for key, value in row.items():
            if key in {"project", "session_id", "model"}:
                continue
            remainder[key] = value
        out.append({**canonical, **remainder})

    return out


def _write_rows_sheet(*, wb: Workbook, mode: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=mode)
    if not rows:
        ws.append(["project", "session_id", "model", "date"])
        return

    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(col) for col in headers])


def _mode_limit(conn: Any, *, mode: str, requested_limit: int | None) -> int:
    if requested_limit is not None:
        return requested_limit
    table = _MODE_TABLES.get(mode)
    if table is None:
        return 100
    rows = query_rows(conn, sql=f"SELECT COUNT(*) AS n FROM {table}", params={})
    if not rows:
        return 1
    total = rows[0].get("n")
    try:
        count = int(total)
    except Exception:
        count = 0
    return max(count, 1)


def export_excel_report(
    conn: Any,
    *,
    req: ExcelRequest,
    output_path: Path,
) -> dict[str, Any]:
    """Run raw-mode queries and save workbook."""
    raw_sections: dict[str, list[dict[str, Any]]] = {}
    all_session_ids: set[str] = set()
    for mode in req.modes:
        rows = run_mode(
            conn,
            mode=mode,
            req=_ModeReq(limit=_mode_limit(conn, mode=mode, requested_limit=req.limit)),
        )
        raw_sections[mode] = rows
        for row in rows:
            sid = _row_session_id(row)
            if sid:
                all_session_ids.add(sid)

    session_dims = _load_session_dims(conn, session_ids=all_session_ids)
    turn_models, latest_models = _load_turn_models(conn, session_ids=all_session_ids)

    filtered_sections: dict[str, list[dict[str, Any]]] = {}
    for mode, rows in raw_sections.items():
        filtered_sections[mode] = _enrich_rows(
            rows=rows,
            req=req,
            session_dims=session_dims,
            turn_models=turn_models,
            latest_models=latest_models,
        )

    wb = Workbook()
    wb.remove(wb.active)
    row_counts: dict[str, int] = {}
    for mode in req.modes:
        rows = filtered_sections.get(mode, [])
        row_counts[mode] = len(rows)
        _write_rows_sheet(wb=wb, mode=mode, rows=rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    return {
        "format": "xlsx",
        "tabs": req.modes,
        "row_counts": row_counts,
        "output_path": str(output_path),
    }
